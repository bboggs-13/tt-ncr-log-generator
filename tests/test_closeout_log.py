import io
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

import engine

CSV = Path("/home/user/134_jane_street_export_nyc_special_inspections___non_conformance_log_30_09_2026.csv")


def _sample_csv() -> bytes:
    rows = []
    for i, status in enumerate(["Open", "Closed", "In Progress", "Closed", "Open", "close"], start=1):
        rows.append({
            "Report #": "65", "Non-Conformance #": f"65.NC.{i}", "Observation Date": "2026-09-01",
            "Type": "TR-1", "Description": f"Item {i}", "Recommended Action": "Fix",
            "Status": status, "Remarks": "", "Date Closed": "2026-09-20" if "clos" in status.lower() else "",
        })
    return pd.DataFrame(rows).to_csv(index=False).encode()


class CloseoutLogFilterTests(unittest.TestCase):
    def _sheet(self, csv_bytes):
        xlsx, df, msgs = engine.process_csv_bytes(csv_bytes, "special_closeout", planradar_access=False)
        self.assertIsNotNone(xlsx, msgs)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))
        return wb, wb.worksheets[0], len(df)

    def test_closed_items_are_hidden_but_kept(self):
        wb, ws, n = self._sheet(_sample_csv())
        start = engine.DATA_START_ROW
        status = {r: ws[f"G{r}"].value for r in range(start, start + n)}
        hidden = {r for r in status if ws.row_dimensions[r].hidden}
        self.assertEqual(len(status), 6)   # every item is still in the log
        self.assertEqual({status[r].lower() for r in hidden}, {"closed", "close"})
        self.assertEqual({status[r] for r in status if r not in hidden}, {"Open", "In Progress"})
        filter_column = ws.auto_filter.filterColumn[0]
        self.assertEqual(filter_column.colId, 6)   # Status column
        self.assertEqual(sorted(filter_column.filters.filter), ["In Progress", "Open"])

    def test_audit_sheet_records_the_default_filter(self):
        wb, ws, n = self._sheet(_sample_csv())
        audit = {a.value: b.value for a, b in wb["Import Audit"].iter_rows(max_col=2)}
        self.assertEqual(audit["Default filter"], "Open items only (3 closed hidden)")


if __name__ == "__main__":
    unittest.main()
