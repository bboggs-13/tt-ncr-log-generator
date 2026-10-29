"""TTFE Inspection Log — import engine.

Turns a PlanRadar CSV export into a standard TT Excel log. Supports multiple
inspection types, each with its own template and column layout but identical
formatting (logo, borders, 10pt font, Status dropdown + filter + live
conditional highlighting, audit sheet).

Add or change a type by editing the PROFILES dict below — no UI changes needed.
"""
from __future__ import annotations

import colorsys
import io
from copy import copy
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl.worksheet.filters import CustomFilter, CustomFilters, FilterColumn
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.formatting.rule import FormulaRule

ASSETS = Path(__file__).parent / "assets"
LOGO_PATH = ASSETS / "tt_logo.png"

# --- Shared layout (same for every template) ---
SHEET = "Sheet1"                     # legacy name (templates are now titled per log type)
SHEET_TITLES = {                     # Excel tab names (31-char limit)
    "special": "NCR Log",
    "special_time": "Time Sensitive NCR Log",
    "special_closeout": "Project Closeout NCR Log",
    "regular": "Site Inspection Log",
    "regular_time": "Time Sensitive Site Insp. Log",
}
SUMMARY_SHEET = "Summary"
AUDIT_SHEET = "Import Audit"
FREEZE_CELL = "A8"                   # rows 1-7 (logo, project, date, legend, column titles) stay visible
HEADER_ROW = 7
DATA_START_ROW = 8
TITLE_ROW = 4
DATE_ROW = 5
FONT_SIZE = 10.0

# Status fills / options / highlighting
OPEN_FILL = PatternFill(start_color="FFBDD7EE", end_color="FFBDD7EE", fill_type="solid")
CLOSED_FILL = PatternFill(start_color="FFC6E0B4", end_color="FFC6E0B4", fill_type="solid")
INPROGRESS_FILL = PatternFill(start_color="FFFFE699", end_color="FFFFE699", fill_type="solid")
STATUS_OPTIONS = ["Open", "In Progress", "Closed"]

# Time Sensitive NCR Log: age bands measured from the log's "Date Updated" (C5).
# (min_days, max_days or None, label, fill) — muted, incrementally stronger tones.
AGE_BANDS = [
    (30, 60, "30–60 days", PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")),
    (60, 90, "60–90 days", PatternFill(start_color="FFF8CBAD", end_color="FFF8CBAD", fill_type="solid")),
    (90, None, "90+ days", PatternFill(start_color="FFF0A8A8", end_color="FFF0A8A8", fill_type="solid")),
]
DATE_FMT = "yyyy-mm-dd"
# Items closed recently: non-overlapping bands of days since Date Closed (inclusive),
# plus a total card covering all of them. (min_days, max_days, label, colour)
CLOSED_BANDS = (
    (0, 29, "Last 30 days", "#C6E0B4"),      # log Closed green
    (30, 59, "30–59 days ago", "#D5E8C8"),
    (60, 90, "60–90 days ago", "#E2EFDA"),
)
CLOSED_TOTAL = (0, 90, "Total, last 90 days", "#A9D08E")
TEXT_FIELDS = ("Description", "Recommended Action", "Remarks")  # left + vertically centered
CLOSED_WORDS = ("closed", "close")


def _pastel(hex6: str, saturation: float = 0.65, lightness: float = 0.86) -> PatternFill:
    """Pastel version of a brand colour: keep its hue, set saturation and lightness."""
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    hue, _l, _s = colorsys.rgb_to_hls(r, g, b)
    t = "".join(f"{round(c * 255):02X}" for c in colorsys.hls_to_rgb(hue, lightness, saturation))
    return PatternFill(start_color="FF" + t, end_color="FF" + t, fill_type="solid")


# Project Closeout NCR Log: rows shaded by NCR type, pastel TT primary colours
# (TR-1 soft green #CDEBCF; pastel Terracotta #D4451D and Dark Blue #004050).
TYPE_FILLS = [
    ("TR-1", "Green", PatternFill(start_color="FFCDEBCF", end_color="FFCDEBCF", fill_type="solid")),  # soft green (not TT olive)
    ("TR-8", "Terracotta", _pastel("D4451D")),
    ("QA/QC", "Dark Blue", _pastel("004050")),
]

# Borders
_THIN = Side(style="thin", color="FF000000")
_THICK = Side(style="thick", color="FF000000")
THIN_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
THICK_BORDER = Border(left=_THICK, right=_THICK, top=_THICK, bottom=_THICK)

LOGO_WIDTH_IN = 3.0
EMU_PER_IN = 914400

# Alternate CSV header names accepted for a given logical field (first match wins).
ALIASES = {
    "Report #": ["Report", "Report No", "Report No.", "Report Number", "Report#"],
    "Non-Conformance #": ["Item #", "Item", "Observation #", "Inspection #", "Issue #"],
    "Type": ["NCR Type", "Type of NCR", "Category"],
    "Responsible Party": ["Responsible", "Party Responsible", "Responsible party",
                          "Trade", "Contractor", "Assigned to"],
}


@dataclass(frozen=True)
class Profile:
    key: str
    label: str
    template: str                 # filename in assets/
    column_map: dict              # csv field -> template column letter
    required_columns: list        # logical field names that must resolve
    center_columns: list          # 1-based column indices to center
    status_col: int               # 1-based index of the Status column
    last_col: int                 # 1-based index of the last column
    age_highlight: bool = False   # Time Sensitive: shade open rows by age of Observation Date
    type_highlight: bool = False  # Project Closeout: shade rows by NCR Type
    status_colors: bool = True    # blue/yellow/green Status-cell colours (off when rows are shaded)
    filter_open: bool = True      # every log opens pre-filtered to non-Closed items
    legend_cols: tuple = ("E", "F")  # legend: (swatch column, label column)
    ui_note: str = ""             # short explanation shown in the app


_AGE_NOTE = ("Open items are shaded across the whole row by days since the Observation Date, "
             "measured from the log's Date Updated: under 30 none · 30–59 yellow · 60–89 orange · "
             "90+ red. Closed items are not shaded.")
_FILTER_NOTE = ("The log opens showing open items only. Closed items are still in the log — "
                "choose Select All in the Status filter to show them.")

# Special Inspections columns (Type and Priority removed)
_SPECIAL_MAP = {
    "Report #": "A", "Non-Conformance #": "B", "Observation Date": "C",
    "Description": "D", "Recommended Action": "E", "Status": "F",
    "Remarks": "G", "Date Closed": "H", "Link to ticket": "I",
}
_SPECIAL_REQ = ["Report #", "Non-Conformance #", "Observation Date", "Status"]
_SPECIAL_CENTER = [1, 2, 3, 6, 8, 9]

_REGULAR_MAP = {
    "Report #": "A", "Non-Conformance #": "B", "Observation Date": "C",
    "Responsible Party": "D", "Description": "E", "Recommended Action": "F",
    "Status": "G", "Remarks": "H", "Date Closed": "I",
}
_REGULAR_REQ = ["Observation Date", "Status"]
_REGULAR_CENTER = [1, 2, 3, 4, 7, 9]


PROFILES = {
    "special": Profile(
        key="special",
        label="Special Inspections",
        template="TTFE_Template.xlsx",
        column_map=_SPECIAL_MAP, required_columns=_SPECIAL_REQ,
        center_columns=_SPECIAL_CENTER, status_col=6, last_col=9,
        ui_note=_FILTER_NOTE,
    ),
    "special_time": Profile(
        key="special_time",
        label="Special Inspections — Time Sensitive NCR Log",
        template="TTFE_Template_TimeSensitive.xlsx",
        column_map=_SPECIAL_MAP, required_columns=_SPECIAL_REQ,
        center_columns=_SPECIAL_CENTER, status_col=6, last_col=9,
        age_highlight=True, status_colors=False, filter_open=True,
        legend_cols=("F", "G"),  # narrow Status column holds the swatches
        ui_note=_AGE_NOTE + " " + _FILTER_NOTE,
    ),
    "special_closeout": Profile(
        key="special_closeout",
        label="Special Inspections — Project Closeout NCR Log",
        template="TTFE_Template_Closeout.xlsx",
        column_map={
            "Report #": "A", "Non-Conformance #": "B", "Observation Date": "C",
            "Type": "D", "Description": "E", "Recommended Action": "F",
            "Status": "G", "Remarks": "H", "Date Closed": "I", "Link to ticket": "J",
        },
        required_columns=["Report #", "Non-Conformance #", "Observation Date", "Type", "Status"],
        center_columns=[1, 2, 3, 4, 7, 9, 10], status_col=7, last_col=10,
        type_highlight=True, status_colors=False, filter_open=True,
        legend_cols=("D", "E"),
        ui_note=("Rows are shaded by NCR Type in pastel colours: TR-1 green · TR-8 terracotta · "
                 "QA/QC dark blue. Other or blank types are not shaded. " + _FILTER_NOTE),
    ),
    "regular": Profile(
        key="regular",
        label="Regular Site Inspections",
        template="TTFE_Template_Regular.xlsx",
        column_map=_REGULAR_MAP, required_columns=_REGULAR_REQ,
        center_columns=_REGULAR_CENTER, status_col=7, last_col=9,
        ui_note=_FILTER_NOTE,
    ),
    "regular_time": Profile(
        key="regular_time",
        label="Regular Site Inspections — Time Sensitive Site Inspection Log",
        template="TTFE_Template_Regular_TimeSensitive.xlsx",
        column_map=_REGULAR_MAP, required_columns=_REGULAR_REQ,
        center_columns=_REGULAR_CENTER, status_col=7, last_col=9,
        age_highlight=True, status_colors=False, filter_open=True,
        legend_cols=("D", "E"),  # narrow Responsible Party column holds the swatches
        ui_note=_AGE_NOTE + " " + _FILTER_NOTE,
    ),
}
DEFAULT_PROFILE = "special"

# What the UI shows: inspection types, and (optionally) log variants within a type.
INSPECTION_TYPES = {"special": "Special Inspections", "regular": "Regular Site Inspections"}
LOG_VARIANTS = {
    "special": {"Standard NCR Log": "special", "Time Sensitive NCR Log": "special_time",
                "Project Closeout NCR Log": "special_closeout"},
    "regular": {"Standard Site Inspection Log": "regular",
                "Time Sensitive Site Inspection Log": "regular_time"},
}


def resolve_profile_key(inspection_type: str, log_variant: str | None = None) -> str:
    variants = LOG_VARIANTS.get(inspection_type)
    if variants and log_variant in variants:
        return variants[log_variant]
    return inspection_type


def get_profile(key: str) -> Profile:
    return PROFILES.get(key, PROFILES[DEFAULT_PROFILE])


# PlanRadar link column: added as the last column when the contractor has PlanRadar
# access, left out when they don't (the link is no use to them).
LINK_FIELD = "Link to ticket"
LINK_WIDTH = 67.86  # same width as the link column in the Special templates


def with_planradar_link(profile: Profile, include: bool) -> Profile:
    """Return the profile with the PlanRadar link column added or removed."""
    has = LINK_FIELD in profile.column_map
    if include and not has:
        col = profile.last_col + 1
        return replace(profile,
                       column_map={**profile.column_map, LINK_FIELD: get_column_letter(col)},
                       center_columns=[*profile.center_columns, col], last_col=col)
    if not include and has:
        col = column_index_from_string(profile.column_map[LINK_FIELD])
        assert col == profile.last_col, "link column must be the last column"
        return replace(profile,
                       column_map={k: v for k, v in profile.column_map.items() if k != LINK_FIELD},
                       center_columns=[c for c in profile.center_columns if c != col],
                       last_col=col - 1)
    return profile


def _fit_link_column(ws, profile: Profile):
    """Make the template's header row match the profile's link setting."""
    link = profile.column_map.get(LINK_FIELD)
    if link:
        head = ws[f"{link}{HEADER_ROW}"]
        if head.value != LINK_FIELD:  # template has no link column: style it like its neighbour
            prev = ws.cell(HEADER_ROW, profile.last_col - 1)
            head.value = LINK_FIELD
            for attr in ("font", "fill", "alignment", "border", "number_format", "protection"):
                setattr(head, attr, copy(getattr(prev, attr)))
        ws.column_dimensions[link].width = LINK_WIDTH
    else:
        col = profile.last_col + 1
        if ws.cell(HEADER_ROW, col).value == LINK_FIELD:  # template has one: drop it
            ws.delete_cols(col)
            ws.column_dimensions.pop(get_column_letter(col), None)


def _resolve_field(field_name: str, df: pd.DataFrame):
    """Return the actual CSV column for a logical field, honoring aliases."""
    if field_name in df.columns:
        return field_name
    for alt in ALIASES.get(field_name, []):
        if alt in df.columns:
            return alt
    return None


def _coerce(val):
    s = str(val).replace(",", "").replace("$", "").strip()
    if s == "":
        return None
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return val


def expected_columns(profile: Profile) -> list:
    return list(profile.column_map.keys())


def validate(df: pd.DataFrame, profile: Profile):
    msgs = []
    missing = [c for c in profile.required_columns if _resolve_field(c, df) is None]
    if missing:
        msgs.append(f"Missing required columns: {', '.join(missing)}")
    if df.empty:
        msgs.append("The CSV has no data rows.")
    mapped = set()
    for fld in profile.column_map:
        actual = _resolve_field(fld, df)
        if actual:
            mapped.add(actual)
    unknown = [c for c in df.columns if c not in mapped and c != LINK_FIELD]
    if unknown:
        msgs.append(f"Note: unmapped columns will be ignored: {', '.join(unknown)}")
    if LINK_FIELD in profile.column_map and _resolve_field(LINK_FIELD, df) is None:
        msgs.append(f"Note: no '{LINK_FIELD}' column in the CSV, so the link column is left blank.")
    ok = not missing and not df.empty
    return ok, msgs


def _apply_age_highlight(ws, profile, n, last_row, last_col_letter):
    """Shade whole data rows by days between Observation Date and Date Updated (C5).

    C5 and the Observation Date cells are written as real dates so Excel can do the
    math; editing C5 later re-ages the log automatically.
    """
    gen = date.today()
    ws["C5"] = gen
    ws["C5"].number_format = DATE_FMT

    obs_letter = profile.column_map["Observation Date"]
    for r in range(DATA_START_ROW, DATA_START_ROW + n):
        cell = ws[f"{obs_letter}{r}"]
        parsed = pd.to_datetime(cell.value, errors="coerce") if cell.value not in (None, "") else pd.NaT
        if not pd.isna(parsed):
            cell.value = parsed.date()
            cell.number_format = DATE_FMT

    dc_letter = profile.column_map.get("Date Closed")
    if dc_letter:
        for r in range(DATA_START_ROW, DATA_START_ROW + n):
            cell = ws[f"{dc_letter}{r}"]
            parsed = pd.to_datetime(cell.value, errors="coerce") if cell.value not in (None, "") else pd.NaT
            if not pd.isna(parsed):
                cell.value = parsed.date()
                cell.number_format = DATE_FMT

    if n > 0:
        rng = f"A{DATA_START_ROW}:{last_col_letter}{last_row}"
        age = f"($C$5-${obs_letter}{DATA_START_ROW})"
        # Only items that are still open get aged; Closed items are never shaded.
        st_cell = f"LOWER(TRIM(${get_column_letter(profile.status_col)}{DATA_START_ROW}))"
        not_closed = ",".join(f'{st_cell}<>"{w}"' for w in CLOSED_WORDS)
        for lo, hi, _label, fill in AGE_BANDS:
            cond = f"{age}>={lo}" if hi is None else f"{age}>={lo},{age}<{hi}"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f"AND(ISNUMBER(${obs_letter}{DATA_START_ROW}),{not_closed},{cond})"], fill=fill))

    rows = [("Less than 30 days (no highlight)", None)] + [
        (f"{lo} to {hi - 1} days" if hi else f"{lo}+ days", fill) for lo, hi, _l, fill in AGE_BANDS
    ]
    _write_legend(ws, profile,
                  "Legend: Days Since Observation Date, Open Items Only (as of Date Updated)", rows)


def _apply_type_highlight(ws, profile, n, last_row, last_col_letter):
    """Shade whole data rows by NCR Type (TR-1 / TR-8 / QA/QC)."""
    t = profile.column_map["Type"]
    if n > 0:
        rng = f"A{DATA_START_ROW}:{last_col_letter}{last_row}"
        for value, _name, fill in TYPE_FILLS:
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'EXACT(UPPER(TRIM(${t}{DATA_START_ROW})),"{value}")'], fill=fill, stopIfTrue=True))
    rows = [(v, fill) for v, _name, fill in TYPE_FILLS] + [("Other / blank type (no highlight)", None)]
    _write_legend(ws, profile, "Legend: NCR Type", rows)


def _apply_open_filter(ws, profile, n, last_row):
    """Open the log filtered to non-Closed items; returns how many closed rows are hidden.

    Uses an Excel custom filter "Status does not equal Closed and does not equal Close",
    so blank / Open / In Progress (or any other status) stay visible, and re-applying the
    filter in Excel after editing a status hides newly closed items too. Closed rows are
    hidden, not deleted: Select All in the Status filter brings them back. Applied even
    when every item is closed, so a download never opens showing closed items.
    """
    letter = get_column_letter(profile.status_col)
    hidden = [
        r for r in range(DATA_START_ROW, last_row + 1)
        if str(ws[f"{letter}{r}"].value or "").strip().lower() in CLOSED_WORDS
    ]
    ws.auto_filter.filterColumn.append(FilterColumn(
        colId=profile.status_col - 1,
        customFilters=CustomFilters(_and=True, customFilter=[
            CustomFilter(operator="notEqual", val=word.capitalize()) for word in CLOSED_WORDS
        ]),
    ))
    for r in hidden:
        ws.row_dimensions[r].hidden = True
    return len(hidden)


def _write_legend(ws, profile, title, rows):
    """Colour key in rows 1-5: swatch in one column, meaning in the next.

    Column pair comes from Profile.legend_cols; fills come from the same constants
    as the highlighting rules, so the key always matches the rows.
    """
    sw_col, lab_col = profile.legend_cols
    head = ws[f"{sw_col}1"]
    head.value = title
    head.font = Font(name="Arial", size=10, bold=True)
    head.alignment = Alignment(horizontal="left", vertical="center")
    head.border = THIN_BORDER
    ws.merge_cells(f"{sw_col}1:{lab_col}1")

    for r, (text, fill) in enumerate(rows, start=2):
        swatch, label = ws[f"{sw_col}{r}"], ws[f"{lab_col}{r}"]
        swatch.border = THIN_BORDER
        if fill is not None:
            swatch.fill = fill
        label.value = text
        label.font = Font(name="Arial", size=10)
        label.alignment = Alignment(horizontal="left", vertical="center")
        label.border = THIN_BORDER


# ---------------------------------------------------------------- summary sheet
_S_GRAY, _S_OLIVE, _S_MUTED, _S_RULE = "FF42424A", "FF839058", "FF6E6E76", "FFC9C9CC"


def _summary_grid(card_px=173, gap_px=18):
    """Column grid where a 4-card row and a 3-card row both span the same width.

    Column edges are the union of both rows' card/gutter edges, so every card is a
    merge of whole columns. Returns (column widths in Excel units, 4-card spans, 3-card spans).
    """
    total = 4 * card_px + 3 * gap_px
    wide = (total - 2 * gap_px) / 3
    def spans(n, w):
        return [(i * (w + gap_px), i * (w + gap_px) + w) for i in range(n)]
    four, three = spans(4, card_px), spans(3, wide)
    edges = sorted({round(e, 2) for a, b in four + three for e in (a, b)})
    widths = [(b - a - 5) / 7 for a, b in zip(edges, edges[1:])]  # px -> Excel char width
    def cols(row):
        out = []
        for a, b in row:
            first = edges.index(round(a, 2)) + 1
            last = edges.index(round(b, 2))
            out.append((get_column_letter(first), get_column_letter(last)))
        return out
    return widths, cols(four), cols(three)


_GRID_WIDTHS, _CARDS4, _CARDS3 = _summary_grid()
_LAST_COL = get_column_letter(len(_GRID_WIDTHS))


def _sfont(size=10, bold=False, color=_S_GRAY):
    return Font(name="Arial", size=size, bold=bold, color=color)


def _section_row(sm, r, title):
    sm[f"A{r}"] = title.upper()
    sm[f"A{r}"].font = _sfont(9, True)
    rule = Side(style="thin", color=_S_GRAY)
    for c in range(1, len(_GRID_WIDTHS) + 1):
        sm.cell(r, c).border = Border(bottom=rule)
    sm.row_dimensions[r].height = 18


def _card_row(sm, top, cards, spans):
    """Four-row cards: colour band, count, label, detail line (each row merged across the card)."""
    edge = Side(style="thin", color=_S_RULE)
    sm.row_dimensions[top].height = 8
    sm.row_dimensions[top + 1].height = 40
    sm.row_dimensions[top + 2].height = 16
    sm.row_dimensions[top + 3].height = 16
    fonts = [None, _sfont(26, True), _sfont(10, True), _sfont(8, color=_S_MUTED)]
    for (c1, c2), card in zip(spans, cards):
        i1, i2 = column_index_from_string(c1), column_index_from_string(c2)
        for k in range(4):
            r = top + k
            if i2 > i1:
                sm.merge_cells(f"{c1}{r}:{c2}{r}")
            for ci in range(i1, i2 + 1):
                cell = sm.cell(r, ci)
                cell.border = Border(left=edge if ci == i1 else None, right=edge if ci == i2 else None,
                                     top=edge if k in (0, 1) else None, bottom=edge if k == 3 else None)
                if k == 0 and card["colour"]:
                    hexc = "FF" + card["colour"].lstrip("#")
                    cell.fill = PatternFill(start_color=hexc, end_color=hexc, fill_type="solid")
            head = sm[f"{c1}{r}"]
            head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if k:
                head.value = (card["count"], card["label"], card["detail"])[k - 1]
                head.font = fonts[k]


def _add_summary_sheet(wb, ws, profile, n, project_name):
    """Summary of the Time Sensitive log, matching the summary PDF.

    Every figure is a formula on the log sheet (same rules as the row highlighting), so
    editing a Status or the Date Updated cell updates the summary as well.
    """
    sm = wb.create_sheet(SUMMARY_SHEET, 0)
    q = "'" + ws.title.replace("'", "''") + "'!"
    last = max(DATA_START_ROW + n - 1, DATA_START_ROW)
    rng = lambda letter: f"{q}${letter}${DATA_START_ROW}:${letter}${last}"
    S = rng(get_column_letter(profile.status_col))
    O = rng(profile.column_map["Observation Date"])
    D = rng(profile.column_map["Date Closed"])
    R = f"{q}$C$5"
    not_closed = f'{S},"<>Closed",{S},"<>Close"'
    log_name = "NCR Log" if profile.key.startswith("special") else "Site Inspection Log"
    kind = "Special Inspections" if profile.key.startswith("special") else "Regular Site Inspections"

    sm.sheet_view.showGridLines = False
    for idx, w in enumerate(_GRID_WIDTHS, start=1):
        sm.column_dimensions[get_column_letter(idx)].width = round(w, 2)

    # Logo, same size and position as the log
    from PIL import Image as PILImage
    w_px, h_px = PILImage.open(LOGO_PATH).size
    width_emu = int(EMU_PER_IN * LOGO_WIDTH_IN)
    img = XLImage(str(LOGO_PATH))
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=0, colOff=0, row=0, rowOff=0),
                               ext=XDRPositiveSize2D(width_emu, int(width_emu * h_px / w_px)))
    sm.add_image(img)

    sm["A4"] = "Time Sensitive NCR Summary"
    sm["A4"].font = _sfont(18, True)
    sm.row_dimensions[4].height = 26
    sm["A5"] = f"{kind}  ·  Open items by days since observation"
    sm["A5"].font = _sfont(11, color=_S_OLIVE)

    # Project / distribution details, aligned to the four card columns
    first = [c1 for c1, _c2 in _CARDS4]
    for col, label in zip(first, ("Project", "Date of distribution", "Open items", "Closed items")):
        sm[f"{col}7"] = label.upper()
        sm[f"{col}7"].font = _sfont(7.5, True, _S_MUTED)
    p_col, d_col, o_col, c_col = first
    sm[f"{p_col}8"] = project_name
    sm[f"{d_col}8"] = f"={R}"
    sm[f"{d_col}8"].number_format = "mmmm d, yyyy"
    sm[f"{c_col}8"] = f'=COUNTIF({S},"Closed")+COUNTIF({S},"Close")'
    sm[f"{o_col}8"] = f"={n}-{c_col}8"
    open_ref = f"${o_col}$8"
    rule = Side(style="thin", color=_S_RULE)
    for ci in range(1, len(_GRID_WIDTHS) + 1):
        cell = sm.cell(8, ci)
        cell.font = _sfont(12)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = Border(bottom=rule)
    for c1, c2 in _CARDS4:            # merge each detail across its card's columns
        if c1 != c2:
            sm.merge_cells(f"{c1}7:{c2}7")
            sm.merge_cells(f"{c1}8:{c2}8")
    sm.row_dimensions[8].height = 22

    # Open items by days since observation (same bands as the log highlighting)
    _section_row(sm, 10, "Open items by days since observation")
    bounds = [("", f'">"&({R}-30)')] + [
        (f'"<="&({R}-{lo})', f'">"&({R}-{hi})' if hi else "") for lo, hi, _l, _f in AGE_BANDS]
    labels = ["Less than 30 days"] + [f"{lo}–{hi - 1} days" if hi else f"{lo}+ days" for lo, hi, _l, _f in AGE_BANDS]
    colours = [None] + ["#" + f.fgColor.rgb[-6:] for _lo, _hi, _l, f in AGE_BANDS]
    details = ["Not highlighted in log", "Highlighted yellow in log",
               "Highlighted orange in log", "Highlighted red in log"]
    cards = []
    for (upper, lower), label, colour, detail in zip(bounds, labels, colours, details):
        crit = "".join(f",{O},{c}" for c in (upper, lower) if c)
        cards.append({"count": f"=COUNTIFS({not_closed}{crit})", "label": label,
                      "colour": colour, "detail": detail})
    _card_row(sm, 12, cards, _CARDS4)
    age_counts = ",".join(f"{c1}13" for c1, _c2 in _CARDS4)

    # Items closed recently: three separate bands plus a 90-day total
    _section_row(sm, 17, "Items closed recently")
    cards = []
    for lo, hi, label, colour in (*CLOSED_BANDS, CLOSED_TOTAL):
        a, b = f"({R}-{hi})", f"({R}-{lo})"
        win = f'{D},">="&{a},{D},"<="&{b}'
        cards.append({
            "count": f'=COUNTIFS({S},"Closed",{win})+COUNTIFS({S},"Close",{win})',
            "label": label, "colour": colour,
            "detail": (f'=TEXT({a},IF(YEAR({a})=YEAR({b}),"mmm d","mmm d, yyyy"))'
                       f'&" – "&TEXT({b},"mmm d, yyyy")'),
        })
    _card_row(sm, 19, cards, _CARDS4)
    sm["A23"] = ("The first three counts do not overlap; the total is their sum. "
                 "Based on each item's Date Closed.")
    sm["A23"].font = _sfont(8, color=_S_MUTED)

    # Notes (live)
    _section_row(sm, 25, "Notes")
    oldest = f'_xlfn.MINIFS({O},{not_closed})'
    notes = [
        f'="•  Ages are counted from each item’s Observation Date to the date of distribution ("'
        f'&TEXT({R},"mmmm d, yyyy")&"), the same rule used to highlight the Time Sensitive {log_name}."',
        "•  Open items include anything not marked Closed (Open and In Progress).",
        f'="•  Closed items are excluded from the age counts: "&{c_col}8&" of {n} items in this log."',
        f'=IF(OR({open_ref}=0,{oldest}=0),"","•  Oldest open item: "&({R}-{oldest})&" days, observed "&TEXT({oldest},"mmmm d, yyyy")&".")',
        f'=IF({open_ref}-SUM({age_counts})>0,"•  "&({open_ref}-SUM({age_counts}))&" open item(s) have no valid Observation Date and cannot be aged.","")',
    ]
    for i, text in enumerate(notes):
        r = 27 + i
        sm.merge_cells(f"A{r}:{_LAST_COL}{r}")
        sm[f"A{r}"] = text
        sm[f"A{r}"].font = _sfont(9.5)
        sm[f"A{r}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        sm.row_dimensions[r].height = 26 if i == 0 else 15

    # Print: one page, portrait
    sm.page_setup.orientation = "portrait"
    sm.page_setup.fitToWidth = 1
    sm.page_setup.fitToHeight = 1
    sm.sheet_properties.pageSetUpPr.fitToPage = True
    sm.print_area = f"A1:{_LAST_COL}32"

    # Open on the summary (and don't leave the log tab grouped with it)
    ws.sheet_view.tabSelected = False
    wb.active = 0


def build_workbook(df: pd.DataFrame, profile: Profile, project_name: str = "134 Jane Street") -> bytes:
    wb = load_workbook(ASSETS / profile.template)
    ws = wb.worksheets[0]
    ws.title = SHEET_TITLES.get(profile.key, "Log")
    ws._images = []  # we place exactly one logo
    _fit_link_column(ws, profile)

    first_col, last_col = 1, profile.last_col
    status_col = profile.status_col

    # Header cells
    ws["A4"] = project_name
    ws["C5"] = date.today().strftime("%Y-%m-%d")

    # Data rows (resolve each field through aliases)
    n = len(df)
    resolved = {fld: _resolve_field(fld, df) for fld in profile.column_map}
    for i, (_, row) in enumerate(df.iterrows()):
        r = DATA_START_ROW + i
        for fld, col_letter in profile.column_map.items():
            src = resolved[fld]
            if src is not None:
                ws[f"{col_letter}{r}"] = _coerce(row[src])

    # Center fixed columns
    for i in range(n):
        r = DATA_START_ROW + i
        for c in profile.center_columns:
            cell = ws.cell(r, c)
            al = copy(cell.alignment)
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=al.wrap_text, text_rotation=al.text_rotation,
                                       indent=al.indent)

    # Text columns: left aligned, vertically centered, wrapped
    text_cols = [profile.column_map[f] for f in TEXT_FIELDS if f in profile.column_map]
    for i in range(n):
        r = DATA_START_ROW + i
        for col_letter in text_cols:
            cell = ws[f"{col_letter}{r}"]
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Borders: thin on all data cells; thick on header row only
    for r in range(DATA_START_ROW, DATA_START_ROW + n):
        for c in range(first_col, last_col + 1):
            ws.cell(r, c).border = THIN_BORDER
    for c in range(first_col, last_col + 1):
        ws.cell(HEADER_ROW, c).border = THICK_BORDER

    # Font: 10pt everywhere (title, headers, data), preserving bold/name/color
    for r in range(TITLE_ROW, DATA_START_ROW + n):
        for c in range(first_col, last_col + 1):
            cell = ws.cell(r, c)
            f = cell.font
            cell.font = Font(name=f.name, size=FONT_SIZE, bold=f.bold,
                             italic=f.italic, color=f.color, underline=f.underline)

    # Filter on the header row
    last_row = DATA_START_ROW + n - 1
    last_col_letter = get_column_letter(last_col)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col_letter}{last_row}"

    # Dropdown on the Status column
    status_letter = get_column_letter(status_col)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"',
                        allow_blank=True, showDropDown=False)
    dv.error = "Choose Open, In Progress, or Closed."
    dv.errorTitle = "Invalid status"
    dv.prompt = "Select a status"
    ws.add_data_validation(dv)
    dv.add(f"{status_letter}{DATA_START_ROW}:{status_letter}{last_row}")

    # Live Status-cell highlighting (recolors on dropdown change). Turned off for logs
    # whose rows are shaded by another rule, so that colour runs across the whole row.
    if profile.status_colors:
        status_range = f"{status_letter}{DATA_START_ROW}:{status_letter}{last_row}"
        top = f"{status_letter}{DATA_START_ROW}"
        ws.conditional_formatting.add(status_range,
            FormulaRule(formula=[f'EXACT(LOWER(TRIM({top})),"open")'], fill=OPEN_FILL, stopIfTrue=True))
        ws.conditional_formatting.add(status_range,
            FormulaRule(formula=[f'EXACT(LOWER(TRIM({top})),"in progress")'], fill=INPROGRESS_FILL, stopIfTrue=True))
        ws.conditional_formatting.add(status_range,
            FormulaRule(formula=[f'EXACT(LOWER(TRIM({top})),"closed")'], fill=CLOSED_FILL, stopIfTrue=True))

    if profile.age_highlight:
        _apply_age_highlight(ws, profile, n, last_row, last_col_letter)
    if profile.type_highlight:
        _apply_type_highlight(ws, profile, n, last_row, last_col_letter)
    hidden_closed = _apply_open_filter(ws, profile, n, last_row) if profile.filter_open and n else 0

    # Logo, exactly 3in wide, top-left corner
    from PIL import Image as PILImage
    w_px, h_px = PILImage.open(LOGO_PATH).size
    width_emu = int(EMU_PER_IN * LOGO_WIDTH_IN)
    height_emu = int(width_emu * h_px / w_px)
    img = XLImage(str(LOGO_PATH))
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=0, colOff=0, row=0, rowOff=0),
                               ext=XDRPositiveSize2D(width_emu, height_emu))
    ws.add_image(img)

    # Keep the title block and column headers visible while scrolling
    ws.freeze_panes = FREEZE_CELL

    # Time Sensitive logs: live summary sheet first, log second
    if profile.age_highlight:
        _add_summary_sheet(wb, ws, profile, n, project_name)

    # Audit sheet
    audit = wb.create_sheet(AUDIT_SHEET)
    for idx, (k, v) in enumerate([
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Inspection type", profile.label),
        ("Aging reference date", ws["C5"].value.strftime("%Y-%m-%d") if profile.age_highlight else "n/a"),
        ("Template", profile.template),
        ("PlanRadar link column", "Included" if LINK_FIELD in profile.column_map else "Omitted (no PlanRadar access)"),
        ("Default filter", f"Open items only ({hidden_closed} closed hidden)" if profile.filter_open and n else "None"),
        ("Rows imported", n),
        ("Project", project_name),
    ], start=1):
        audit[f"A{idx}"] = k
        audit[f"B{idx}"] = v

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def process_csv_bytes(csv_bytes: bytes, inspection_type: str = DEFAULT_PROFILE,
                      project_name: str = "134 Jane Street", planradar_access: bool = True):
    """Full pipeline. Returns (xlsx_bytes_or_None, df, messages).

    planradar_access: True adds the PlanRadar "Link to ticket" column, False leaves it out.
    """
    profile = with_planradar_link(get_profile(inspection_type), planradar_access)
    df = pd.read_csv(io.BytesIO(csv_bytes), dtype=str).fillna("")
    ok, msgs = validate(df, profile)
    if not ok:
        return None, df, msgs
    xlsx = build_workbook(df, profile, project_name=project_name)
    return xlsx, df, msgs
